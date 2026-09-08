"""
duplicate_finder.py
-------------------
Encontra ficheiros duplicados em uma ou mais diretorias (e sub-diretorias),
gera um relatório Excel e move os duplicados para uma pasta de quarentena.

Modos de utilização:

  1) ANÁLISE — gera o relatório Excel (com coluna editável "Manter?"):
        python duplicate_finder.py <dir1> [dir2 ...] [--quarantine <pasta>] [--move]

  2) APLICAR — lê um Excel já preenchido e move os não-mantidos para quarentena:
        python duplicate_finder.py --apply <relatorio.xlsx> [--quarantine <pasta>] [--dry-run]

Fluxo recomendado (revisão humana):
    a) Corre a análise (sem --move) para gerar o Excel.
    b) Abre o Excel, folha "Detalhes", e marca SIM na coluna "Manter?" nas
       cópias que queres preservar (pelo menos uma por grupo; podes manter várias).
    c) Corre com --apply <excel> para mover as restantes para quarentena.
       Cada grupo é validado: se algum ficar sem nenhum SIM, nada é movido.

Exemplos:
    python duplicate_finder.py "C:\\Projetos\\Cliente"
    python duplicate_finder.py "C:\\Pasta1" "C:\\Pasta2"
    python duplicate_finder.py "C:\\Pasta1" "C:\\Pasta2" --quarantine "C:\\Quarentena" --move
    python duplicate_finder.py --apply "C:\\Pasta1\\relatorio_duplicados_20250101_120000.xlsx"
    python duplicate_finder.py --apply "C:\\...\\relatorio.xlsx" --dry-run

Dependências (instalar uma vez):
    pip install openpyxl
"""

import os
import sys
import hashlib

# Garante que os emojis/acentos não rebentam na consola Windows (cp1252)
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass
import shutil
import argparse
from collections import defaultdict
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ── Utilitários ──────────────────────────────────────────────────────────────

def format_size(size_bytes):
    """Converte bytes para formato legível."""
    for unit in ["B", "KB", "MB", "GB", "TB"]:
        if size_bytes < 1024:
            return f"{size_bytes:.2f} {unit}"
        size_bytes /= 1024
    return f"{size_bytes:.2f} PB"


def hash_file(path, chunk_size=8192):
    """Calcula o hash MD5 de um ficheiro."""
    h = hashlib.md5()
    try:
        with open(path, "rb") as f:
            while chunk := f.read(chunk_size):
                h.update(chunk)
        return h.hexdigest()
    except (PermissionError, OSError):
        return None


# ── Lógica principal ─────────────────────────────────────────────────────────

def scan_files(root_dirs):
    """Percorre uma ou mais diretorias e agrupa todos os ficheiros por tamanho."""
    size_map = defaultdict(list)
    total = 0

    for root_dir in root_dirs:
        print(f"\n🔍 A analisar: {root_dir}")
        dir_total = 0
        for dirpath, _, filenames in os.walk(root_dir):
            for fname in filenames:
                fpath = os.path.join(dirpath, fname)
                try:
                    size = os.path.getsize(fpath)
                    size_map[size].append(fpath)
                    dir_total += 1
                except (PermissionError, OSError):
                    continue
        print(f"   ✅ {dir_total} ficheiros encontrados")
        total += dir_total

    print(f"\n📁 Total combinado: {total} ficheiro(s) em {len(root_dirs)} diretoria(s)")
    return size_map


def find_duplicates(size_map):
    """
    Agrupa por tamanho primeiro (rápido), depois calcula hash
    apenas nos grupos com mais de 1 ficheiro (eficiente).
    """
    print("⚙️  A calcular hashes dos candidatos a duplicados...")
    hash_map = defaultdict(list)
    candidates = sum(len(v) for v in size_map.values() if len(v) > 1)
    processed = 0

    for paths in size_map.values():
        if len(paths) < 2:
            continue
        for path in paths:
            fhash = hash_file(path)
            if fhash:
                hash_map[fhash].append(path)
            processed += 1
            if processed % 50 == 0:
                print(f"   ... {processed}/{candidates} ficheiros processados")

    duplicates = {h: ps for h, ps in hash_map.items() if len(ps) > 1}
    total_dup_groups = len(duplicates)
    total_dup_files = sum(len(v) - 1 for v in duplicates.values())
    print(f"   ✅ {total_dup_groups} grupos de duplicados | {total_dup_files} ficheiros a remover")
    return duplicates


# ── Relatório Excel ──────────────────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", start_color="1F4E79")
GROUP_FILL    = PatternFill("solid", start_color="D6E4F0")
ORIGINAL_FILL = PatternFill("solid", start_color="E2EFDA")
DUP_FILL      = PatternFill("solid", start_color="FCE4D6")
EDIT_FILL     = PatternFill("solid", start_color="FFF2CC")  # coluna "Manter?" (editável)

# Valores aceites na coluna "Manter?" como "manter este ficheiro"
KEEP_VALUES = {"SIM", "X", "YES", "TRUE", "1", "MANTER", "V", "✓"}
HEADER_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT     = Font(name="Arial", size=10)
BOLD_FONT     = Font(name="Arial", bold=True, size=10)
CENTER        = Alignment(horizontal="center", vertical="center")
LEFT          = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)


def style_cell(cell, font=None, fill=None, alignment=None, border=None):
    if font:      cell.font      = font
    if fill:      cell.fill      = fill
    if alignment: cell.alignment = alignment
    if border:    cell.border    = border


def create_excel_report(duplicates, output_path, moved_files):
    """Gera o relatório Excel com duas folhas: Resumo e Detalhes."""
    wb = Workbook()

    # ── Folha 1: Resumo ───────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Resumo"
    ws_sum.sheet_view.showGridLines = False

    # Título
    ws_sum.merge_cells("A1:D1")
    ws_sum["A1"] = "Relatório de Ficheiros Duplicados"
    style_cell(ws_sum["A1"], font=Font(name="Arial", bold=True, size=16, color="1F4E79"), alignment=CENTER)
    ws_sum.row_dimensions[1].height = 35

    ws_sum.merge_cells("A2:D2")
    ws_sum["A2"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    style_cell(ws_sum["A2"], font=Font(name="Arial", size=10, color="595959"), alignment=CENTER)
    ws_sum.row_dimensions[2].height = 20

    ws_sum.append([])  # linha vazia

    # Métricas
    total_groups   = len(duplicates)
    total_dup_files = sum(len(v) - 1 for v in duplicates.values())
    total_wasted   = sum(os.path.getsize(p) for paths in duplicates.values() for p in paths[1:] if os.path.exists(p))

    headers_sum = ["Métrica", "Valor"]
    ws_sum.append(headers_sum)
    for col, h in enumerate(headers_sum, 1):
        cell = ws_sum.cell(row=4, column=col)
        style_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=thin_border)

    metrics = [
        ("Grupos de duplicados encontrados", total_groups),
        ("Ficheiros duplicados (a remover)", total_dup_files),
        ("Espaço desperdiçado",              format_size(total_wasted)),
        ("Ficheiros movidos para quarentena", len(moved_files)),
    ]
    for i, (label, value) in enumerate(metrics, 5):
        ws_sum[f"A{i}"] = label
        ws_sum[f"B{i}"] = value
        style_cell(ws_sum[f"A{i}"], font=BODY_FONT, alignment=LEFT,   border=thin_border)
        style_cell(ws_sum[f"B{i}"], font=BOLD_FONT, alignment=CENTER, border=thin_border)

    ws_sum.column_dimensions["A"].width = 42
    ws_sum.column_dimensions["B"].width = 30

    # ── Folha 2: Detalhes ─────────────────────────────────────────────────
    ws = wb.create_sheet("Detalhes")
    ws.sheet_view.showGridLines = False

    col_headers = ["Grupo", "Manter?", "Tipo", "Nome do Ficheiro", "Caminho Completo",
                   "Tamanho", "Data de Modificação", "Estado", "Hash"]
    ws.append(col_headers)

    for col, h in enumerate(col_headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = h
        style_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=CENTER, border=thin_border)
    ws.row_dimensions[1].height = 22

    # Índices de coluna (1-based)
    COL_MANTER = 2
    COL_HASH   = 9

    # Dropdown SIM/NÃO na coluna "Manter?"
    dv = DataValidation(type="list", formula1='"SIM,NÃO"', allow_blank=True)
    dv.prompt = "Escreve SIM para manter este ficheiro; deixa vazio (ou NÃO) para mover."
    dv.promptTitle = "Manter ficheiro?"
    ws.add_data_validation(dv)

    row = 2
    for group_num, (fhash, paths) in enumerate(sorted(duplicates.items()), 1):
        for idx, path in enumerate(paths):
            tipo  = "✅ Original" if idx == 0 else "🔁 Duplicado"
            fill  = ORIGINAL_FILL if idx == 0 else DUP_FILL
            fname = os.path.basename(path)
            # Pré-preenche "SIM" na cópia auto-detetada como original
            manter = "SIM" if idx == 0 else ""

            try:
                size  = os.path.getsize(path) if os.path.exists(path) else 0
                mtime = datetime.fromtimestamp(os.path.getmtime(path)).strftime("%d/%m/%Y %H:%M") \
                        if os.path.exists(path) else "—"
            except OSError:
                size, mtime = 0, "—"

            estado = "Movido para quarentena" if path in moved_files else (
                     "Original (mantido)"     if idx == 0 else "Duplicado (não movido)")

            values = [group_num, manter, tipo, fname, path,
                      format_size(size), mtime, estado, fhash]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                # A coluna "Manter?" fica com fill amarelo para sinalizar que é editável
                cell_fill = EDIT_FILL if col == COL_MANTER else fill
                align = CENTER if col == COL_MANTER else LEFT
                style_cell(cell, font=BODY_FONT, fill=cell_fill, alignment=align, border=thin_border)

            dv.add(ws.cell(row=row, column=COL_MANTER))
            ws.row_dimensions[row].height = 18
            row += 1

        # Linha separadora entre grupos
        ws.append([])
        row += 1

    # Larguras das colunas
    col_widths = [8, 10, 14, 35, 65, 12, 22, 28, 34]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freezar cabeçalho
    ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"\n📊 Relatório Excel guardado em: {output_path}")


# ── Quarentena ───────────────────────────────────────────────────────────────

def move_to_quarantine(duplicates, quarantine_dir):
    """Move os duplicados (mantém sempre o primeiro) para a pasta de quarentena."""
    os.makedirs(quarantine_dir, exist_ok=True)
    moved = []

    print(f"\n📦 A mover duplicados para: {quarantine_dir}")
    for paths in duplicates.values():
        for path in paths[1:]:  # O índice 0 é o "original" — fica no lugar
            try:
                dest_name = os.path.basename(path)
                dest      = os.path.join(quarantine_dir, dest_name)
                # Evita colisões de nome na quarentena
                if os.path.exists(dest):
                    base, ext = os.path.splitext(dest_name)
                    dest = os.path.join(quarantine_dir, f"{base}_{len(moved)}{ext}")
                shutil.move(path, dest)
                moved.append(path)
                print(f"   ↪ {os.path.basename(path)}")
            except (PermissionError, OSError) as e:
                print(f"   ⚠️  Erro ao mover {path}: {e}")

    print(f"   ✅ {len(moved)} ficheiro(s) movido(s)")
    return moved


# ── Aplicar decisões a partir do Excel ───────────────────────────────────────

def _is_keep(value):
    """Interpreta o valor da coluna 'Manter?' como decisão de manter."""
    if value is None:
        return False
    return str(value).strip().upper() in KEEP_VALUES


def read_decisions(excel_path):
    """
    Lê a folha 'Detalhes' do Excel e devolve os grupos com as decisões.
    Estrutura devolvida: { hash: [ {path, keep, group, row}, ... ] }
    """
    wb = load_workbook(excel_path, data_only=True)
    if "Detalhes" not in wb.sheetnames:
        raise ValueError("O Excel não tem uma folha 'Detalhes'. "
                         "Usa um relatório gerado por este script.")
    ws = wb["Detalhes"]

    # Mapeia cabeçalhos -> índice de coluna (1-based), tolerante a reordenação
    header = {}
    for col, cell in enumerate(ws[1], 1):
        if cell.value:
            header[str(cell.value).strip().lower()] = col

    required = ["caminho completo", "hash", "manter?", "grupo"]
    missing = [h for h in required if h not in header]
    if missing:
        raise ValueError(f"Colunas em falta no Excel: {', '.join(missing)}. "
                         "O ficheiro tem de ser um relatório gerado por este script.")

    c_path  = header["caminho completo"]
    c_hash  = header["hash"]
    c_keep  = header["manter?"]
    c_group = header["grupo"]

    groups = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        path = ws.cell(row=r, column=c_path).value
        fhash = ws.cell(row=r, column=c_hash).value
        if not path or not fhash:  # linha separadora ou vazia
            continue
        groups[str(fhash)].append({
            "path":  str(path),
            "keep":  _is_keep(ws.cell(row=r, column=c_keep).value),
            "group": ws.cell(row=r, column=c_group).value,
            "row":   r,
        })
    return groups


def validate_decisions(groups):
    """
    Garante que cada grupo tem PELO MENOS um ficheiro marcado para manter.
    Devolve a lista de grupos inválidos (vazia se tudo ok).
    """
    invalid = []
    for fhash, items in groups.items():
        if not any(it["keep"] for it in items):
            invalid.append((fhash, items))
    return invalid


def apply_from_excel(excel_path, quarantine_dir, dry_run=False):
    """Lê o Excel, valida as decisões e move os não-mantidos para quarentena."""
    print(f"\n📖 A ler decisões de: {excel_path}")
    groups = read_decisions(excel_path)

    total_files = sum(len(v) for v in groups.values())
    print(f"   ✅ {len(groups)} grupo(s) | {total_files} ficheiro(s) no relatório")

    # 1. Validação — nenhum grupo pode ficar sem "manter"
    invalid = validate_decisions(groups)
    if invalid:
        print(f"\n❌ Validação falhou: {len(invalid)} grupo(s) sem nenhum ficheiro marcado como 'manter':")
        for fhash, items in invalid:
            grp = items[0]["group"] if items else "?"
            print(f"   • Grupo {grp} (hash {fhash[:8]}…) — {len(items)} cópias, 0 marcadas com SIM")
            for it in items:
                print(f"       - {it['path']}")
        print("\n   Nenhum ficheiro foi movido. Corrige o Excel (marca SIM em pelo menos "
              "uma linha de cada grupo) e volta a correr.")
        return "invalid", []

    # 2. Recolhe o que vai ser movido (tudo o que não está marcado para manter)
    to_move = []
    for items in groups.values():
        for it in items:
            if not it["keep"]:
                to_move.append(it["path"])

    kept = total_files - len(to_move)
    print(f"\n📋 Resumo: manter {kept} | mover {len(to_move)}")

    if dry_run:
        print("\nℹ️  Modo simulação (--dry-run) — nada foi movido. Ficheiros que seriam movidos:")
        for p in to_move:
            print(f"   ↪ {p}")
        return "dry_run", []

    # 3. Mover para quarentena
    os.makedirs(quarantine_dir, exist_ok=True)
    moved, errors = [], []
    print(f"\n📦 A mover não-mantidos para: {quarantine_dir}")
    for path in to_move:
        if not os.path.exists(path):
            print(f"   ⚠️  Já não existe (ignorado): {path}")
            continue
        try:
            dest_name = os.path.basename(path)
            dest = os.path.join(quarantine_dir, dest_name)
            if os.path.exists(dest):
                base, ext = os.path.splitext(dest_name)
                dest = os.path.join(quarantine_dir, f"{base}_{len(moved)}{ext}")
            shutil.move(path, dest)
            moved.append(path)
            print(f"   ↪ {dest_name}")
        except (PermissionError, OSError) as e:
            errors.append((path, str(e)))
            print(f"   ⚠️  Erro ao mover {path}: {e}")

    print(f"\n   ✅ {len(moved)} ficheiro(s) movido(s)" +
          (f" | ⚠️  {len(errors)} erro(s)" if errors else ""))
    return "ok", moved


# ── Entry point ───────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Detetor de ficheiros duplicados com relatório Excel e quarentena."
    )
    parser.add_argument("roots", nargs="*",
        help="Uma ou mais diretorias a analisar. Ex: C:\\Pasta1 C:\\Pasta2")
    parser.add_argument("--quarantine", default=None,
        help="Pasta de quarentena. Por omissão: <primeira diretoria>\\Quarentena_Duplicados")
    parser.add_argument("--move", action="store_true",
        help="Ativa a movimentação dos duplicados para quarentena. "
             "Sem esta flag só gera o relatório (modo seguro).")
    parser.add_argument("--apply", metavar="EXCEL", default=None,
        help="Lê um relatório Excel já preenchido (coluna 'Manter?') e move os "
             "não-mantidos para quarentena. Valida que cada grupo tem pelo menos "
             "um ficheiro marcado com SIM.")
    parser.add_argument("--dry-run", action="store_true",
        help="Com --apply: valida e mostra o que seria movido, sem mover nada.")

    args = parser.parse_args()

    # ── Modo APLICAR: executa a partir de um Excel preenchido ─────────────
    if args.apply:
        excel_path = os.path.abspath(args.apply)
        if not os.path.isfile(excel_path):
            print(f"❌ Excel não encontrado: {excel_path}")
            sys.exit(1)
        quarantine_dir = args.quarantine or os.path.join(
            os.path.dirname(excel_path), "Quarentena_Duplicados")
        try:
            status, moved = apply_from_excel(excel_path, quarantine_dir, dry_run=args.dry_run)
        except ValueError as e:
            print(f"❌ {e}")
            sys.exit(1)

        if status == "invalid":
            sys.exit(2)  # validação falhou — nada movido
        print("\n✅ Concluído!")
        if status == "ok":
            print(f"   Quarentena : {quarantine_dir}\n")
        sys.exit(0)

    # ── Modo ANÁLISE ──────────────────────────────────────────────────────
    if not args.roots:
        print("❌ Indica pelo menos uma diretoria a analisar, ou usa --apply <excel>.")
        parser.print_help()
        sys.exit(1)

    roots = [os.path.abspath(r) for r in args.roots]

    for r in roots:
        if not os.path.isdir(r):
            print(f"❌ Diretoria não encontrada: {r}")
            sys.exit(1)

    quarantine_dir = args.quarantine or os.path.join(roots[0], "Quarentena_Duplicados")
    timestamp      = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path    = os.path.join(roots[0], f"relatorio_duplicados_{timestamp}.xlsx")

    # 1. Scan
    size_map   = scan_files(roots)

    # 2. Duplicados
    duplicates = find_duplicates(size_map)

    if not duplicates:
        print("\n✅ Nenhum ficheiro duplicado encontrado!")
        sys.exit(0)

    # 3. Mover (opcional)
    moved_files = []
    if args.move:
        moved_files = move_to_quarantine(duplicates, quarantine_dir)
    else:
        print("\nℹ️  Modo de leitura — nenhum ficheiro foi movido.")
        print("   Para mover os duplicados, adiciona --move ao comando.")

    # 4. Relatório Excel
    create_excel_report(duplicates, report_path, moved_files)

    print("\n✅ Concluído!")
    if args.move:
        print(f"   Quarentena : {quarantine_dir}")
    print(f"   Relatório  : {report_path}\n")


if __name__ == "__main__":
    main()
